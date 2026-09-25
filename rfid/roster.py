"""Чтение и заполнение списков групп, которые даёт пользователь.

Источник истины по составу групп — сами файлы, а не наша база. Программа
их не создаёт: она находит, разбирает и дописывает в них колонки с датами.

Предмет — это **имя папки** внутри tables/:

    tables/
      Бургеростроение 1 курс/        <- предмет
        1000000.10001 список.xlsx    <- одна группа
        1000000.10002 список.xlsx    <- другая группа

Пример структуры списка с вымышленным номером группы:

    A1  «Группа 1000000/10001», объединено A1:Y1
    A2  «№», B2 «ФИО студента», C2:Y2 объединено и пусто
    A3+ номера — ФОРМУЛЫ вида =A2+1, а не числа
    B3+ фамилии
    C28:Y28 объединённое примечание «Староста: …» под списком

Поэтому: номера берём по порядку строк, а не из ячеек; конец списка ищем
по пустому ФИО; перед записью дат объединение в строке заголовка снимаем.
"""

from __future__ import annotations

import os
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

DATE_HEADER_FORMAT = "%d.%m"
ABSENT_MARK = "—"
LATE_AFTER = timedelta(minutes=10)
LATE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF00")

# Расписание со слов пользователя: пара 1 час 40 минут, перемена 20 минут,
# то есть занятия начинаются ровно каждые два часа.
LESSON_LENGTH = timedelta(hours=1, minutes=40)
BREAK_LENGTH = timedelta(minutes=20)

# По циклу приходы режутся на занятия: marks приходит общий на предмет,
# а у групп пара бывает в разное время — без разделения дневная группа
# целиком красилась бы как опоздавшая. Общая лекция двух групп при этом
# остаётся одним занятием, и опоздавшая группа на ней видна.
#
# Порогом служит ВЕСЬ цикл, а не одна пара: иначе карта, приложенная
# на перемене, создала бы фантомное занятие, и следующая пара целиком
# оказалась бы опоздавшей. Приход на перемене лучше отнести к предыдущей
# паре, чем испортить следующую.
LESSON_CYCLE = LESSON_LENGTH + BREAK_LENGTH

# Сколько строк сверху просматривать в поисках шапки.
_HEADER_SEARCH_DEPTH = 15

_GROUP_RE = re.compile(r"группа\s*[:№]?\s*(?P<name>\S.*)", re.IGNORECASE)
_NAME_HEADER = re.compile(r"ФИО|фамилия", re.IGNORECASE)

# Номер группы «по виду»: 1000000/10001, 1000000-10001 и подобные.
# Нужен, когда слова «Группа» в файле нет — например «ИКНК 1000000/10002».
_GROUP_CODE_RE = re.compile(r"\d{3,}\s*[/\\.\-]\s*\d{2,}")


class RosterError(Exception):
    """Файл не похож на список группы."""


@dataclass(frozen=True, slots=True)
class RosterStudent:
    number: int      # порядковый номер в списке
    full_name: str
    row: int         # строка на листе


@dataclass(slots=True)
class Roster:
    path: Path
    sheet_title: str
    group_name: str
    header_row: int
    name_col: int
    first_date_col: int
    students: list[RosterStudent] = field(default_factory=list)

    def find(self, full_name: str) -> RosterStudent | None:
        target = _normalize(full_name)
        for student in self.students:
            if _normalize(student.full_name) == target:
                return student
        return None


def normalize_name(name: str) -> str:
    """Для сравнения ФИО: регистр, ё и лишние пробелы значения не имеют."""
    return " ".join(name.replace("ё", "е").replace("Ё", "Е").split()).casefold()


_normalize = normalize_name  # короткий псевдоним для внутреннего кода


# ------------------------------------------------------------------- чтение


def _cells_above_header(sheet: Worksheet, header_row: int):
    for row in range(1, header_row):
        for col in range(1, min(sheet.max_column, 10) + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                yield value.strip()


def _detect_group(sheet: Worksheet, header_row: int, fallback: str) -> str:
    """Название группы из шапки файла.

    Три попытки по убыванию надёжности: явное «Группа X»; номер группы
    по виду (например, «ИКНК 1000000/10002» — без слова
    «Группа»); наконец первая непустая строка над шапкой. Имя файла —
    последнее средство, оно хуже всего описывает группу.
    """
    above = list(_cells_above_header(sheet, header_row))

    for text in above:
        match = _GROUP_RE.match(text)
        if match:
            return match.group("name").strip()

    for text in above:
        match = _GROUP_CODE_RE.search(text)
        if match:
            return re.sub(r"\s*", "", match.group(0))

    return above[0] if above else fallback


def _detect_header(sheet: Worksheet) -> tuple[int, int]:
    """Найти строку шапки и колонку с ФИО."""
    depth = min(sheet.max_row, _HEADER_SEARCH_DEPTH)
    for row in range(1, depth + 1):
        for col in range(1, min(sheet.max_column, 10) + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and _NAME_HEADER.search(value):
                return row, col
    raise RosterError("не нашёл шапку: нет ячейки со словом «ФИО»")


def _read_students(sheet: Worksheet, header_row: int, name_col: int) -> list[RosterStudent]:
    students: list[RosterStudent] = []
    row = header_row + 1
    while row <= sheet.max_row:
        value = sheet.cell(row=row, column=name_col).value
        if value is None or not str(value).strip():
            break  # список кончился; ниже могут быть примечания
        students.append(
            RosterStudent(number=len(students) + 1, full_name=str(value).strip(), row=row)
        )
        row += 1
    return students


def read_roster(path: str | Path) -> Roster:
    path = Path(path)
    workbook = load_workbook(path)
    sheet = workbook.active

    header_row, name_col = _detect_header(sheet)
    students = _read_students(sheet, header_row, name_col)
    if not students:
        raise RosterError(f"в {path.name} не нашлось ни одного студента")

    return Roster(
        path=path,
        sheet_title=sheet.title,
        group_name=_detect_group(sheet, header_row, fallback=path.stem),
        header_row=header_row,
        name_col=name_col,
        first_date_col=name_col + 1,
        students=students,
    )


# -------------------------------------------------------------------- запись


def _unmerge_around(sheet: Worksheet, row: int, from_col: int) -> None:
    """Снять объединения, мешающие писать в строку заголовка дат.

    В настоящем списке C2:Y2 объединено — без этого дату туда не положить.
    """
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row <= row <= merged.max_row and merged.max_col >= from_col:
            sheet.unmerge_cells(str(merged))


def _date_columns(sheet: Worksheet, roster: Roster) -> dict[str, int]:
    """Какие даты уже есть в шапке и в каких колонках."""
    found: dict[str, int] = {}
    for col in range(roster.first_date_col, sheet.max_column + 1):
        value = sheet.cell(row=roster.header_row, column=col).value
        if isinstance(value, str) and value.strip():
            found[value.strip()] = col
        elif isinstance(value, datetime):
            found[value.strftime(DATE_HEADER_FORMAT)] = col
    return found


def _next_free_column(sheet: Worksheet, roster: Roster, taken: dict[str, int]) -> int:
    return max([*taken.values(), roster.first_date_col - 1]) + 1


def session_starts(arrivals: Iterable[datetime]) -> list[datetime]:
    """Начала занятий за день.

    Новое занятие начинается, когда приход отстоит от начала текущего
    на LESSON_CYCLE или больше. Именно от начала, а не от предыдущего
    прихода: опоздавший на час не должен объявлять себя началом новой пары.

    Сравнение нестрогое, и это существенно: пары идут ровно через цикл,
    поэтому приход точно через два часа — это уже следующая пара.
    """
    starts: list[datetime] = []
    current: datetime | None = None
    for at in sorted(arrivals):
        if current is None or at - current >= LESSON_CYCLE:
            starts.append(at)
            current = at
    return starts


def session_start_for(at: datetime, starts: list[datetime]) -> datetime | None:
    """Начало того занятия, к которому относится этот приход."""
    found = None
    for start in starts:
        if start > at:
            break
        found = start
    return found


def is_late(at: datetime, starts: list[datetime]) -> bool:
    start = session_start_for(at, starts)
    return start is not None and at - start >= LATE_AFTER


def write_attendance(
    roster: Roster,
    marks: dict[date, dict[str, datetime]],
    *,
    absent_mark: str = ABSENT_MARK,
) -> Path:
    """Проставить отметки в файле группы, сохранив его оформление.

    marks: дата -> {ФИО: время прихода}. ФИО сверяются без учёта регистра и «ё».
    Возвращает путь файла. Имена, которых нет в списке, молча не теряются —
    их возвращает сверка через unmatched_names().
    """
    workbook = load_workbook(roster.path)
    sheet = workbook[roster.sheet_title]

    _unmerge_around(sheet, roster.header_row, roster.first_date_col)
    existing = _date_columns(sheet, roster)

    for day in sorted(marks):
        title = day.strftime(DATE_HEADER_FORMAT)
        column = existing.get(title)
        if column is None:
            column = _next_free_column(sheet, roster, existing)
            existing[title] = column
            header = sheet.cell(row=roster.header_row, column=column, value=title)
            header.alignment = Alignment(horizontal="center")
            header.font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(column)].width = 7

        # marks содержит все группы предмета: отсчёт общий для предмета за день.
        # marks приходит общий на предмет, поэтому приходы сначала разбиваются
        # на занятия по разрыву во времени — см. session_starts().
        starts = session_starts(marks[day].values())
        present = {_normalize(name): at for name, at in marks[day].items()}
        for student in roster.students:
            at = present.get(_normalize(student.full_name))
            # Например, ФИО в B17:C17 занимает будущую ячейку отметки C17.
            # Освобождаем только нужную ячейку; значение в B17 сохраняется.
            for merged in list(sheet.merged_cells.ranges):
                if (merged.min_row <= student.row <= merged.max_row
                        and merged.min_col <= column <= merged.max_col):
                    sheet.unmerge_cells(str(merged))
            cell = sheet.cell(row=student.row, column=column)
            if at is None and cell_time(cell.value) is not None:
                # В файле что-то стоит, а в базе отметки нет: время, вписанное
                # руками, или непонятное «+». Файл главнее — не затираем.
                continue
            cell.value = at.strftime("%H:%M") if at else absent_mark
            cell.alignment = Alignment(horizontal="center")
            if at is not None and is_late(at, starts):
                cell.fill = LATE_FILL
            elif cell.fill == LATE_FILL:
                # Повторный экспорт должен снимать устаревшую подсветку.
                cell.fill = PatternFill()

    _save_atomically(workbook, roster.path)
    return roster.path


def _save_atomically(workbook, path: Path) -> None:
    """Сохранить через временный файл рядом, потом подменить.

    Прямой workbook.save(path) сначала обнуляет файл и только потом пишет:
    сорвавшаяся запись — блокировка, падение, нет места на диске — оставляет
    от списка группы пустышку. А список дал пользователь, восстановить его
    неоткуда: в базе лежат отметки, а не состав групп.

    Здесь оригинал не трогается, пока новый файл не собран целиком.
    """
    temp = path.with_name(path.name + ".tmp")
    try:
        workbook.save(temp)
        os.replace(temp, path)
    except BaseException:
        # Не оставляем мусор рядом со списком, каким бы ни был сбой.
        try:
            temp.unlink(missing_ok=True)
        except OSError:
            pass
        raise


# ------------------------------------------------------ чтение отметок из файла

_TIME_RE = re.compile(r"^(?P<h>\d{1,2})[:.](?P<m>\d{2})(?::\d{2})?$")
_DAY_RE = re.compile(r"^(?P<d>\d{1,2})\.(?P<m>\d{1,2})(?:\.(?P<y>\d{2}|\d{4}))?$")


@dataclass(frozen=True, slots=True)
class FileMarks:
    """Отметки, найденные в файле группы."""

    # дата -> [(студент, время прихода)]
    marks: dict[date, list[tuple[RosterStudent, datetime]]]
    # Что не удалось понять: (заголовок колонки или ФИО, значение).
    # Такое не угадываем, а показываем человеку.
    skipped: list[tuple[str, str]]
    # дата -> кто в этот день явно не был (пусто или прочерк)
    absent: dict[date, list[RosterStudent]] = field(default_factory=dict)


def header_date(value, today: date) -> date | None:
    """Дата из заголовка колонки.

    Мы пишем «дд.мм» без года. Год восстанавливается так, чтобы дата
    не оказалась в будущем: «12.09», прочитанное в январе, — прошлый сентябрь.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None
    match = _DAY_RE.match(value.strip())
    if not match:
        return None
    day, month = int(match["d"]), int(match["m"])
    year = match["y"]
    try:
        if year:
            return date(int(year) + (2000 if len(year) == 2 else 0), month, day)
        guess = date(today.year, month, day)
        return guess if guess <= today else date(today.year - 1, month, day)
    except ValueError:
        return None


def cell_time(value) -> time | None | str:
    """Время прихода из ячейки.

    time — пришёл; None — пусто или прочерк, то есть не был; строка —
    непонятное значение, которое нельзя молча принять ни за приход,
    ни за пропуск.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = str(value).strip()
    if not text or text in {ABSENT_MARK, "-", "–", "н", "Н", "нб", "НБ"}:
        return None
    match = _TIME_RE.match(text)
    if match and int(match["h"]) < 24 and int(match["m"]) < 60:
        return time(int(match["h"]), int(match["m"]))
    return text


def read_attendance(roster: Roster, today: date | None = None) -> FileMarks:
    """Прочитать уже проставленные в файле отметки."""
    today = today or date.today()
    sheet = load_workbook(roster.path)[roster.sheet_title]

    marks: dict[date, list[tuple[RosterStudent, datetime]]] = {}
    absent: dict[date, list[RosterStudent]] = {}
    skipped: list[tuple[str, str]] = []
    for col in range(roster.first_date_col, sheet.max_column + 1):
        head = sheet.cell(row=roster.header_row, column=col).value
        day = header_date(head, today)
        if day is None:
            continue  # не дата — чужая колонка, например «Примечание»
        for student in roster.students:
            value = cell_time(sheet.cell(row=student.row, column=col).value)
            if value is None:
                absent.setdefault(day, []).append(student)
                continue
            if isinstance(value, str):
                skipped.append((f"{head}, {student.full_name}", value))
                continue
            marks.setdefault(day, []).append((student, datetime.combine(day, value)))
    return FileMarks(marks, skipped, absent)


def unmatched_names(roster: Roster, names: list[str]) -> list[str]:
    """Кого из отмеченных не нашлось в списке группы.

    Связка «карта → студент» живёт в базе и опирается на ФИО, поэтому
    правка фамилии в файле рвёт связь. Молчать об этом нельзя.
    """
    return [name for name in names if roster.find(name) is None]


# ------------------------------------------------------- предметы = папки


@dataclass(frozen=True, slots=True)
class SubjectFolder:
    """Предмет — это папка. Её имя и есть название предмета с курсом."""

    name: str
    path: Path
    rosters: tuple[Path, ...]

    @property
    def group_count(self) -> int:
        return len(self.rosters)


def _roster_files(folder: Path) -> tuple[Path, ...]:
    return tuple(
        sorted(
            p for p in folder.glob("*.xlsx")
            if not p.name.startswith("~$")  # временные файлы открытого Excel
        )
    )


def as_folder(path: str | Path) -> SubjectFolder:
    """Считать конкретную папку предметом: её имя — название, .xlsx внутри — группы."""
    path = Path(path)
    return SubjectFolder(name=path.name, path=path, rosters=_roster_files(path))


def discover_subjects(tables_dir: str | Path) -> list[SubjectFolder]:
    """Предметы — это подпапки tables/.

    Сама tables/ предметом не является: она корень, в котором эти папки лежат.
    Подпапки без .xlsx не показываем — выбирать там нечего.
    """
    root = Path(tables_dir)
    if not root.is_dir():
        return []
    return [
        as_folder(folder)
        for folder in sorted(root.iterdir())
        if folder.is_dir() and _roster_files(folder)
    ]


def misplaced_rosters(tables_dir: str | Path) -> tuple[Path, ...]:
    """Списки, лежащие в корне tables/ мимо папок-предметов.

    Они не попадут ни в один предмет. Сообщаем об этом в диагностике,
    но не трогаем: раскладку ведёт пользователь.
    """
    root = Path(tables_dir)
    return _roster_files(root) if root.is_dir() else ()
