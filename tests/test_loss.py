"""Проверки против потери отметки — самое важное свойство программы.

Студент приложил карту и ушёл. Если отметка не записалась, узнать об этом
уже неоткуда: карта у студента, а очередь идёт дальше. Поэтому здесь
собраны сценарии, в которых отметка могла бы пропасть молча.
"""

import sqlite3
from datetime import datetime

import pytest

from rfid import pipeline
from rfid.codes import parse_line
from rfid.readers import MockCardReader
from rfid.storage import Storage

LINES = [
    "Em-Marine[A100] 007,42",
    "Em-Marine[B200] 008,43",
    "Em-Marine[C300] 009,44",
]


class QuietView:
    def __init__(self):
        self.messages = []

    def status(self, message):
        self.messages.append(message)

    def show(self, result):
        pass


class ViewThatBreaks(QuietView):
    """Вывод падает на второй карте: сбой консоли, звука, чего угодно."""

    def __init__(self, fail_on=2):
        super().__init__()
        self.fail_on = fail_on
        self.shown = 0

    def show(self, result):
        self.shown += 1
        if self.shown == self.fail_on:
            raise RuntimeError("сбой вывода")


@pytest.fixture
def db(tmp_path):
    with Storage(tmp_path / "t.db") as s:
        yield s


def marks_in(storage):
    return storage.conn.execute("SELECT COUNT(*) FROM attendance").fetchone()[0]


class TestLoopSurvives:
    def test_display_failure_does_not_lose_the_rest_of_the_queue(self, db):
        """Сбой показа не имеет права оборвать пару.

        Раньше исключение из view.show убивало цикл, и все последующие
        студенты молча не отмечались.
        """
        subject = db.add_subject("Тест")
        view = ViewThatBreaks()
        tally = pipeline.run(MockCardReader(lines=LINES), db, view, subject)

        assert marks_in(db) == 3, "должны записаться все три карты"
        assert tally[pipeline.SHOW_FAILED] == 1

    def test_display_failure_is_reported(self, db):
        subject = db.add_subject("Тест")
        view = ViewThatBreaks()
        pipeline.run(MockCardReader(lines=LINES), db, view, subject)
        assert any("сбой вывода" in m for m in view.messages)

    def test_write_failure_does_not_lose_the_rest(self, db, monkeypatch):
        """Одна неудачная запись не должна ронять остальные."""
        subject = db.add_subject("Тест")
        calls = {"n": 0}
        real_mark = db.mark

        def flaky(code, subj, **kw):
            calls["n"] += 1
            if calls["n"] == 2:
                raise sqlite3.OperationalError("database is locked")
            return real_mark(code, subj, **kw)

        monkeypatch.setattr(db, "mark", flaky)
        view = QuietView()
        tally = pipeline.run(MockCardReader(lines=LINES), db, view, subject)

        assert marks_in(db) == 2, "первая и третья должны сохраниться"
        assert tally[pipeline.FAILED] == 1

    def test_write_failure_is_shouted_about(self, db, monkeypatch):
        """Потерянная отметка обязана быть видна, а не утонуть в логе."""
        subject = db.add_subject("Тест")

        def always_fails(*_a, **_kw):
            raise sqlite3.OperationalError("database is locked")

        monkeypatch.setattr(db, "mark", always_fails)
        view = QuietView()
        tally = pipeline.run(MockCardReader(lines=LINES[:1]), db, view, subject)

        assert pipeline.failed_count(tally) == 1
        assert any("НЕ ЗАПИСАНО" in m for m in view.messages)
        assert any("A10007002A" in m for m in view.messages)


class TestBusyDatabase:
    """Второе окно программы не должно стоить отметок."""

    def test_opening_a_busy_database_works(self, tmp_path):
        path = tmp_path / "t.db"
        with Storage(path) as first:
            first.add_subject("Тест")

        holder = sqlite3.connect(path, isolation_level=None)
        holder.execute("BEGIN IMMEDIATE")
        holder.execute(
            "INSERT INTO subjects(name, course, created_at) VALUES('x','','2026-01-01')"
        )
        try:
            holder.rollback()  # отпускаем, но соединение живо
            with Storage(path) as second:
                assert second.find_subject("Тест") is not None
        finally:
            holder.close()

    def test_busy_timeout_is_set(self, db):
        value = db.conn.execute("PRAGMA busy_timeout").fetchone()[0]
        assert value >= 1000, "мгновенный отказ стоил бы отметки"

    def test_opening_does_not_write_when_nothing_changed(self, tmp_path):
        """Открытие базы не должно брать блокировку записи на пустом месте."""
        path = tmp_path / "t.db"
        with Storage(path) as s:
            s.add_subject("Тест")

        holder = sqlite3.connect(path, isolation_level=None)
        holder.execute("BEGIN IMMEDIATE")   # держим запись
        try:
            with Storage(path) as s:        # не должно упасть
                assert s.list_subjects()
        finally:
            holder.rollback()
            holder.close()


class TestMarkIsCommittedImmediately:
    def test_mark_survives_a_crash(self, tmp_path):
        """Отметка должна быть на диске сразу, а не в конце сеанса."""
        path = tmp_path / "t.db"
        with Storage(path) as s:
            subject = s.add_subject("Тест")
            s.mark(parse_line(LINES[0]), subject, at=datetime(2026, 9, 25, 9, 0))
            # Соединение не закрываем: смотрим файл со стороны, как после сбоя.
            other = sqlite3.connect(path)
            try:
                n = other.execute("SELECT COUNT(*) FROM attendance").fetchone()[0]
            finally:
                other.close()
        assert n == 1, "отметка обязана быть зафиксирована сразу"


class TestExportSurvives:
    """Сбой на одном файле не должен лишать записи остальные группы."""

    @staticmethod
    def _folder_with_bad_file(tmp_path):
        from test_roster import make_roster_file

        folder = tmp_path / "Предмет"
        folder.mkdir()
        (folder / "1-битый.xlsx").write_bytes(b"")          # 0 байт
        (folder / "2-чужой.xlsx").write_bytes(b"not a zip")  # мусор
        make_roster_file(folder / "3-хороший.xlsx",
                         names=["Тестов Тест"], group="1000000/10001")
        return folder

    def test_broken_file_does_not_block_the_others(self, tmp_path, db):
        from openpyxl import load_workbook
        from rfid import journal, roster as R

        folder = self._folder_with_bad_file(tmp_path)
        subject = db.add_subject("Предмет")
        card = parse_line(LINES[0])
        db.add_student(card, "Тестов Тест", "1000000/10001")
        db.mark(card, subject, at=datetime(2026, 9, 25, 9, 0))

        results = journal.fill_subject(db, subject, R.discover_subjects(tmp_path)[0])

        good = [r for r in results if r.ok]
        assert len(good) == 1, "хорошая группа обязана быть заполнена"
        sheet = load_workbook(folder / "3-хороший.xlsx").active
        assert sheet.cell(row=3, column=3).value == "09:00"

    def test_broken_files_are_reported(self, tmp_path, db):
        from rfid import journal, roster as R

        self._folder_with_bad_file(tmp_path)
        subject = db.add_subject("Предмет")
        results = journal.fill_subject(db, subject, R.discover_subjects(tmp_path)[0])

        broken = [r for r in results if not r.ok]
        assert len(broken) == 2, "про каждый сбойный файл надо сказать"
        assert all(r.error for r in broken)
