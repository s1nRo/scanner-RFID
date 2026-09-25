"""Тесты чтения и заполнения списков групп.

Образец воспроизводит формат списка с вымышленными данными: объединённые ячейки,
номера-формулы в колонке A и примечание «Староста» под списком. Именно
на этих особенностях наивный разбор и ломается.
"""

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from rfid import roster as R

NAMES = [
    "Тестов Тест Тестович",
    "Примеров Пример Примерович",
    "Образцова Проба Тестовна",
    "Макетов Пётр Тестович",
]

DAY_ONE = date(2026, 9, 20)
DAY_TWO = date(2026, 9, 22)


def make_roster_file(path: Path, names=NAMES, group="1000000/10001") -> Path:
    """Копия формата пользователя со всеми его неудобствами."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"

    sheet["A1"] = f"Группа {group}"
    sheet.merge_cells("A1:Y1")

    sheet["A2"] = "№"
    sheet["B2"] = "ФИО студента"
    sheet.merge_cells("C2:Y2")  # пустая объединённая шапка под даты

    for index, name in enumerate(names):
        row = 3 + index
        # Номера именно формулами, как в настоящем файле.
        sheet.cell(row=row, column=1, value=1 if index == 0 else f"=A{row - 1}+1")
        sheet.cell(row=row, column=2, value=name)

    note_row = 3 + len(names)
    sheet.cell(row=note_row, column=3, value="Староста: Тестовый Студент; контакт скрыт")
    sheet.merge_cells(start_row=note_row, start_column=3, end_row=note_row, end_column=25)

    workbook.save(path)
    return path


@pytest.fixture
def roster_file(tmp_path):
    return make_roster_file(tmp_path / "группа.xlsx")


class TestReading:
    def test_group_from_title_row(self, roster_file):
        assert R.read_roster(roster_file).group_name == "1000000/10001"

    def test_finds_all_students(self, roster_file):
        students = R.read_roster(roster_file).students
        assert [s.full_name for s in students] == NAMES

    def test_numbering_is_positional_not_from_formulas(self, roster_file):
        """В колонке A формулы — брать номер оттуда нельзя."""
        students = R.read_roster(roster_file).students
        assert [s.number for s in students] == [1, 2, 3, 4]

    def test_note_below_list_is_not_a_student(self, roster_file):
        """«Староста: …» лежит ниже списка и в студенты попасть не должен."""
        names = [s.full_name for s in R.read_roster(roster_file).students]
        assert not any("Староста" in n for n in names)

    def test_rows_point_at_real_cells(self, roster_file):
        sheet = load_workbook(roster_file).active
        for student in R.read_roster(roster_file).students:
            assert sheet.cell(row=student.row, column=2).value == student.full_name

    def test_dates_start_right_after_names(self, roster_file):
        parsed = R.read_roster(roster_file)
        assert parsed.name_col == 2
        assert parsed.first_date_col == 3

    def test_group_code_found_without_the_word(self, tmp_path):
        """Пример «ИКНК 1000000/10002» без слова «Группа»."""
        path = make_roster_file(tmp_path / "Список_10002.xlsx")
        workbook = load_workbook(path)
        workbook.active["A1"] = "ИКНК 1000000/10002"
        workbook.save(path)
        assert R.read_roster(path).group_name == "1000000/10002"

    def test_plain_text_above_header_is_used(self, tmp_path):
        path = make_roster_file(tmp_path / "файл.xlsx")
        workbook = load_workbook(path)
        workbook.active["A1"] = "Первый курс, поток А"
        workbook.save(path)
        assert R.read_roster(path).group_name == "Первый курс, поток А"

    def test_group_falls_back_to_filename(self, tmp_path):
        path = tmp_path / "БИ-11.xlsx"
        make_roster_file(path, group="")
        workbook = load_workbook(path)
        workbook.active["A1"] = None
        workbook.save(path)
        assert R.read_roster(path).group_name == "БИ-11"

    def test_file_without_header_is_rejected(self, tmp_path):
        path = tmp_path / "мусор.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "что-то не то"
        workbook.save(path)
        with pytest.raises(R.RosterError):
            R.read_roster(path)


class TestNameMatching:
    def test_case_and_yo_are_ignored(self, roster_file):
        parsed = R.read_roster(roster_file)
        assert parsed.find("макетов пётр тестович") is not None
        assert parsed.find("МАКЕТОВ ПЕТР ТЕСТОВИЧ") is not None

    def test_extra_spaces_ignored(self, roster_file):
        parsed = R.read_roster(roster_file)
        assert parsed.find("  Образцова   Проба  Тестовна ") is not None

    def test_unknown_name_reported(self, roster_file):
        parsed = R.read_roster(roster_file)
        assert R.unmatched_names(parsed, ["Образцова Проба Тестовна", "Кто-то Левый"]) == [
            "Кто-то Левый"
        ]


class TestWriting:
    def test_merged_name_does_not_block_attendance(self, tmp_path):
        names = [f"Студент {i}" for i in range(15)]
        path = make_roster_file(tmp_path / "merged.xlsx", names=names)
        workbook = load_workbook(path)
        sheet = workbook.active
        sheet.merge_cells("B17:C17")
        sheet.merge_cells("D17:E17")
        workbook.save(path)

        parsed = R.read_roster(path)
        marks = {DAY_ONE: {
            names[0]: datetime(2026, 9, 20, 9, 0),
            names[-1]: datetime(2026, 9, 20, 9, 10),
        }}
        R.write_attendance(parsed, marks)
        R.write_attendance(R.read_roster(path), marks)

        sheet = load_workbook(path).active
        assert sheet["B17"].value == names[-1]
        assert sheet["A17"].value == "=A16+1"
        assert sheet["C17"].value == "09:10"
        assert sheet["C17"].fill.fgColor.rgb == "FFFFFF00"
        assert "B17:C17" not in sheet.merged_cells
        assert "D17:E17" in sheet.merged_cells
        assert "A1:Y1" in sheet.merged_cells
        assert "C18:Y18" in sheet.merged_cells
        assert [s.full_name for s in R.read_roster(path).students] == names

    def test_late_threshold_includes_exactly_ten_minutes(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {
            NAMES[0]: datetime(2026, 9, 20, 9, 0, 30),
            NAMES[1]: datetime(2026, 9, 20, 9, 10, 29),
            NAMES[2]: datetime(2026, 9, 20, 9, 10, 30),
            NAMES[3]: datetime(2026, 9, 20, 9, 15),
        }})
        sheet = load_workbook(roster_file).active
        assert sheet["C3"].fill.patternType is None
        assert sheet["C4"].fill.patternType is None
        for address in ("C5", "C6"):
            assert sheet[address].fill.patternType == "solid"
            assert sheet[address].fill.fgColor.rgb == "FFFFFF00"

    def test_late_threshold_is_independent_for_each_day(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {
            DAY_ONE: {
                NAMES[0]: datetime(2026, 9, 20, 9, 0),
                NAMES[1]: datetime(2026, 9, 20, 9, 10),
            },
            DAY_TWO: {NAMES[1]: datetime(2026, 9, 22, 12, 0)},
        })
        sheet = load_workbook(roster_file).active
        assert sheet["C4"].fill.patternType == "solid"
        assert sheet["D4"].fill.patternType is None
        assert sheet["C5"].fill.patternType is None  # отсутствующий

    def test_other_lesson_the_same_day_is_not_late(self, tmp_path):
        """marks общий на предмет, но у групп пара бывает в разное время.

        Без разделения по разрыву дневная группа целиком красилась бы жёлтым,
        хотя пришла на своё занятие вовремя.
        """
        morning = make_roster_file(tmp_path / "утро.xlsx", names=["Утренний Алексей"])
        afternoon = make_roster_file(tmp_path / "день.xlsx", names=["Дневной Борис"])
        marks = {DAY_ONE: {
            "Утренний Алексей": datetime(2026, 9, 20, 9, 0),
            "Дневной Борис": datetime(2026, 9, 20, 14, 0),
        }}
        for path in (morning, afternoon):
            R.write_attendance(R.read_roster(path), marks)

        for path in (morning, afternoon):
            assert load_workbook(path).active["C3"].fill.patternType is None, path.name

    def test_late_to_a_shared_lesson_is_still_caught(self, tmp_path):
        """Обе группы на одной лекции — опоздавшая обязана быть видна."""
        first = make_roster_file(tmp_path / "а.xlsx", names=["Вовремя Анна"])
        second = make_roster_file(tmp_path / "б.xlsx", names=["Опоздавший Борис"])
        marks = {DAY_ONE: {
            "Вовремя Анна": datetime(2026, 9, 20, 9, 0),
            "Опоздавший Борис": datetime(2026, 9, 20, 9, 20),
        }}
        for path in (first, second):
            R.write_attendance(R.read_roster(path), marks)

        assert load_workbook(first).active["C3"].fill.patternType is None
        assert load_workbook(second).active["C3"].fill.patternType == "solid"

    @pytest.mark.parametrize("still_present", [True, False])
    def test_reexport_clears_outdated_late_fill(self, roster_file, still_present):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {
            NAMES[0]: datetime(2026, 9, 20, 9, 0),
            NAMES[1]: datetime(2026, 9, 20, 9, 10),
        }})
        updated = {NAMES[1]: datetime(2026, 9, 20, 9, 10)} if still_present else {}
        R.write_attendance(parsed, {DAY_ONE: updated})
        sheet = load_workbook(roster_file).active
        assert sheet["C4"].fill.patternType is None
        assert sheet["C4"].value == ("09:10" if still_present else R.ABSENT_MARK)

    def test_date_column_added(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)}})

        sheet = load_workbook(roster_file).active
        assert sheet.cell(row=2, column=3).value == "20.09"
        assert sheet.cell(row=3, column=3).value == "09:02"

    def test_absent_get_a_dash(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)}})
        sheet = load_workbook(roster_file).active
        assert sheet.cell(row=4, column=3).value == R.ABSENT_MARK

    def test_second_date_goes_to_next_column(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {
            DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)},
            DAY_TWO: {NAMES[1]: datetime(2026, 9, 22, 9, 5)},
        })
        sheet = load_workbook(roster_file).active
        assert sheet.cell(row=2, column=3).value == "20.09"
        assert sheet.cell(row=2, column=4).value == "22.09"

    def test_rewriting_same_date_reuses_column(self, roster_file):
        """Повторный экспорт не должен плодить колонки за один и тот же день."""
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)}})
        R.write_attendance(parsed, {DAY_ONE: {NAMES[1]: datetime(2026, 9, 20, 9, 7)}})

        sheet = load_workbook(roster_file).active
        assert sheet.cell(row=2, column=3).value == "20.09"
        assert sheet.cell(row=2, column=4).value is None
        assert sheet.cell(row=3, column=3).value == R.ABSENT_MARK
        assert sheet.cell(row=4, column=3).value == "09:07"

    def test_names_and_numbers_are_not_touched(self, roster_file):
        before = load_workbook(roster_file).active
        names_before = [before.cell(row=r, column=2).value for r in range(3, 7)]

        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)}})

        after = load_workbook(roster_file).active
        assert [after.cell(row=r, column=2).value for r in range(3, 7)] == names_before

    def test_note_under_the_list_survives(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {}})
        sheet = load_workbook(roster_file).active
        note = sheet.cell(row=3 + len(NAMES), column=3).value
        assert note and "Староста" in note

    def test_merged_header_is_released(self, roster_file):
        """C2:Y2 объединено — без снятия дату туда не записать."""
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {}})
        sheet = load_workbook(roster_file).active
        assert "C2:Y2" not in [str(r) for r in sheet.merged_cells.ranges]

    def test_reparse_after_write_is_stable(self, roster_file):
        parsed = R.read_roster(roster_file)
        R.write_attendance(parsed, {DAY_ONE: {NAMES[0]: datetime(2026, 9, 20, 9, 2)}})
        again = R.read_roster(roster_file)
        assert [s.full_name for s in again.students] == NAMES
        assert again.group_name == "1000000/10001"


class TestSubjectFolders:
    def test_folders_with_lists_are_subjects(self, tmp_path):
        for name in ("Бургеростроение 1 курс", "Матанализ 2 курс"):
            (tmp_path / name).mkdir()
            make_roster_file(tmp_path / name / "г.xlsx")
        assert [s.name for s in R.discover_subjects(tmp_path)] == [
            "Бургеростроение 1 курс",
            "Матанализ 2 курс",
        ]

    def test_empty_folders_are_skipped(self, tmp_path):
        """Выбирать в папке без .xlsx нечего — не показываем её."""
        (tmp_path / "Пустая").mkdir()
        assert R.discover_subjects(tmp_path) == []

    def test_root_is_not_a_subject(self, tmp_path):
        """tables — корень, где лежат папки предметов, а не предмет сам."""
        make_roster_file(tmp_path / "россыпью.xlsx")
        assert R.discover_subjects(tmp_path) == []

    def test_only_subfolders_are_returned(self, tmp_path):
        make_roster_file(tmp_path / "россыпью.xlsx")
        (tmp_path / "Физика").mkdir()
        make_roster_file(tmp_path / "Физика" / "г.xlsx")
        assert [f.name for f in R.discover_subjects(tmp_path)] == ["Физика"]

    def test_misplaced_files_are_listed_separately(self, tmp_path):
        make_roster_file(tmp_path / "забытый.xlsx")
        (tmp_path / "Физика").mkdir()
        make_roster_file(tmp_path / "Физика" / "г.xlsx")
        assert [p.name for p in R.misplaced_rosters(tmp_path)] == ["забытый.xlsx"]

    def test_rosters_inside_are_listed(self, tmp_path):
        folder = tmp_path / "Физика 1 курс"
        folder.mkdir()
        make_roster_file(folder / "группа1.xlsx")
        make_roster_file(folder / "группа2.xlsx")
        subject = R.discover_subjects(tmp_path)[0]
        assert subject.group_count == 2

    def test_excel_lock_files_ignored(self, tmp_path):
        folder = tmp_path / "Физика"
        folder.mkdir()
        make_roster_file(folder / "группа.xlsx")
        (folder / "~$группа.xlsx").write_bytes(b"lock")
        assert R.discover_subjects(tmp_path)[0].group_count == 1

    def test_missing_directory_is_not_an_error(self, tmp_path):
        assert R.discover_subjects(tmp_path / "нет") == []

    def test_as_folder_takes_any_directory(self, tmp_path):
        make_roster_file(tmp_path / "г.xlsx")
        folder = R.as_folder(tmp_path)
        assert folder.name == tmp_path.name
        assert folder.group_count == 1


@pytest.fixture
def full_roster_file(tmp_path):
    names = [f"Тестовый Студент {i:02d}" for i in range(1, 26)]
    return make_roster_file(tmp_path / "полный_тестовый_список.xlsx", names=names)


class TestFullRoster:
    """Полный список генерируется без зависимости от личных таблиц."""

    def test_parses(self, full_roster_file):
        parsed = R.read_roster(full_roster_file)
        assert parsed.group_name == "1000000/10001"
        assert len(parsed.students) == 25
        assert parsed.students[0].full_name == "Тестовый Студент 01"
        assert parsed.students[-1].full_name == "Тестовый Студент 25"

    def test_note_not_counted(self, full_roster_file):
        names = [s.full_name for s in R.read_roster(full_roster_file).students]
        assert not any("Староста" in n for n in names)


class TestSessionSplitting:
    """Разбиение приходов на занятия по разрыву во времени."""

    @staticmethod
    def at(hour, minute=0):
        return datetime(2026, 9, 20, hour, minute)

    def test_one_lesson_gives_one_start(self):
        arrivals = [self.at(9, 0), self.at(9, 5), self.at(9, 20)]
        assert R.session_starts(arrivals) == [self.at(9, 0)]

    def test_next_lesson_opens_a_new_session(self):
        arrivals = [self.at(9, 0), self.at(9, 5), self.at(14, 0), self.at(14, 3)]
        assert R.session_starts(arrivals) == [self.at(9, 0), self.at(14, 0)]

    def test_back_to_back_lessons_are_separated(self):
        """Пара 9:00-10:40, перемена 20 минут, следующая ровно в 11:00."""
        arrivals = [self.at(9, 0), self.at(9, 6), self.at(11, 0), self.at(11, 5)]
        assert R.session_starts(arrivals) == [self.at(9, 0), self.at(11, 0)]

    def test_exactly_one_cycle_later_is_a_new_lesson(self):
        """Пары идут ровно через цикл — сравнение обязано быть нестрогим."""
        assert R.session_starts([self.at(9, 0), self.at(11, 0)]) == [
            self.at(9, 0), self.at(11, 0)
        ]

    def test_tap_during_the_break_does_not_poison_the_next_lesson(self):
        """Карта, приложенная на перемене, не должна создавать занятие.

        Иначе следующая пара целиком считалась бы опоздавшей.
        """
        arrivals = [self.at(9, 0), self.at(10, 45), self.at(11, 0), self.at(11, 6)]
        starts = R.session_starts(arrivals)
        assert starts == [self.at(9, 0), self.at(11, 0)]
        assert R.is_late(self.at(11, 0), starts) is False
        assert R.is_late(self.at(11, 6), starts) is False

    def test_order_of_input_does_not_matter(self):
        jumbled = [self.at(14, 3), self.at(9, 0), self.at(14, 0), self.at(9, 5)]
        assert R.session_starts(jumbled) == [self.at(9, 0), self.at(14, 0)]

    def test_no_arrivals(self):
        assert R.session_starts([]) == []

    def test_latecomer_does_not_open_a_new_lesson(self):
        """Иначе опоздавший сам себя объявил бы началом занятия."""
        arrivals = [self.at(9, 0), self.at(9, 40)]
        assert R.session_starts(arrivals) == [self.at(9, 0)]
        assert R.is_late(self.at(9, 40), R.session_starts(arrivals))

    def test_very_late_arrival_still_counts_as_late(self):
        """Пара идёт 1:40 — пришедший через час двадцать всё ещё опоздал."""
        arrivals = [self.at(9, 0), self.at(10, 20)]
        assert R.session_starts(arrivals) == [self.at(9, 0)]
        assert R.is_late(self.at(10, 20), R.session_starts(arrivals))

    def test_arrivals_do_not_drift_into_one_long_session(self):
        """Отсчёт от начала пары, а не от предыдущего прихода.

        Иначе цепочка приходов по чуть-чуть растянула бы одно занятие
        на весь день.
        """
        arrivals = [self.at(9, 0), self.at(10, 0), self.at(11, 0), self.at(12, 0)]
        assert R.session_starts(arrivals) == [self.at(9, 0), self.at(11, 0)]

    def test_cycle_is_longer_than_the_late_window(self):
        """Иначе разделение занятий съело бы саму подсветку опозданий."""
        assert R.LESSON_CYCLE > R.LATE_AFTER

    def test_cycle_is_lesson_plus_break(self):
        assert R.LESSON_CYCLE == R.LESSON_LENGTH + R.BREAK_LENGTH

    def test_arrival_is_matched_to_its_own_lesson(self):
        starts = R.session_starts([self.at(9, 0), self.at(14, 0)])
        assert R.session_start_for(self.at(9, 30), starts) == self.at(9, 0)
        assert R.session_start_for(self.at(14, 30), starts) == self.at(14, 0)

    def test_arrival_before_any_lesson_has_no_start(self):
        starts = R.session_starts([self.at(9, 0)])
        assert R.session_start_for(self.at(8, 0), starts) is None
        assert R.is_late(self.at(8, 0), starts) is False

    def test_late_only_after_the_threshold(self):
        starts = R.session_starts([self.at(9, 0)])
        assert R.is_late(self.at(9, 0) + R.LATE_AFTER, starts) is True
        assert R.is_late(self.at(9, 9), starts) is False


class TestWriteSafety:
    """Список групп даёт пользователь, восстановить его неоткуда.

    Реальная потеря данных 25.09.2026: openpyxl.save() сначала обнуляет файл
    и только потом пишет, поэтому сорвавшаяся запись оставляла от списка
    пустышку в 0 байт. Запись теперь идёт через временный файл.
    """

    def test_failed_write_leaves_the_roster_intact(self, roster_file, monkeypatch):
        before = roster_file.read_bytes()
        assert before[:2] == b"PK"

        def boom(*_a, **_kw):
            raise PermissionError(13, "file is locked")

        monkeypatch.setattr(R.os, "replace", boom)
        with pytest.raises(PermissionError):
            R.write_attendance(R.read_roster(roster_file), {DAY_ONE: {}})

        assert roster_file.read_bytes() == before, "оригинал не должен пострадать"

    def test_broken_save_leaves_the_roster_intact(self, roster_file, monkeypatch):
        """Сбой на середине самой записи тоже не должен трогать оригинал."""
        before = roster_file.read_bytes()

        import openpyxl.workbook.workbook as wb_module

        def boom(self, *_a, **_kw):
            raise OSError("disk full")

        monkeypatch.setattr(wb_module.Workbook, "save", boom)
        with pytest.raises(OSError):
            R.write_attendance(R.read_roster(roster_file), {DAY_ONE: {}})

        assert roster_file.read_bytes() == before

    def test_no_temp_file_left_behind(self, roster_file, monkeypatch):
        def boom(*_a, **_kw):
            raise PermissionError(13, "file is locked")

        monkeypatch.setattr(R.os, "replace", boom)
        with pytest.raises(PermissionError):
            R.write_attendance(R.read_roster(roster_file), {DAY_ONE: {}})

        assert list(roster_file.parent.glob("*.tmp")) == []

    def test_successful_write_leaves_no_temp(self, roster_file):
        R.write_attendance(R.read_roster(roster_file), {DAY_ONE: {}})
        assert list(roster_file.parent.glob("*.tmp")) == []
        assert roster_file.read_bytes()[:2] == b"PK"
