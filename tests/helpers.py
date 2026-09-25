"""Общее для тестов: образец списка группы и помощники без None.

Помощники не только укорачивают тесты. Функции вроде parse_line или
find_subject по контракту могут вернуть None, и тест, который тут же
обращается к результату, упал бы с невнятным AttributeError. Здесь
падение сразу говорит, чего не нашлось.
"""

from __future__ import annotations

from pathlib import Path
from typing import TypeVar

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from rfid.scanner import CardCode, parse_line

T = TypeVar("T")

ROSTER_NAMES = [
    "Тестов Тест Тестович",
    "Примеров Пример Примерович",
    "Образцова Проба Тестовна",
    "Макетов Пётр Тестович",
]


def must(value: T | None) -> T:
    """Значение, которое обязано быть."""
    assert value is not None, "ожидалось значение, получено None"
    return value


def card(line: str) -> CardCode:
    """Код карты из строки считывателя, которая точно является кодом."""
    code = parse_line(line)
    assert code is not None, f"не код карты: {line!r}"
    return code


def active(workbook: Workbook) -> Worksheet:
    """Активный лист книги — обязательно лист с ячейками."""
    sheet = workbook.active
    assert isinstance(sheet, Worksheet), "активный лист не таблица"
    return sheet


def open_sheet(path: Path) -> Worksheet:
    """Открыть файл и взять его активный лист."""
    return active(load_workbook(path))


def make_roster_file(path: Path, names=ROSTER_NAMES, group="1000000/10001") -> Path:
    """Копия формата пользователя со всеми его неудобствами."""
    workbook = Workbook()
    page = active(workbook)
    page.title = "Лист1"

    page["A1"] = f"Группа {group}"
    page.merge_cells("A1:Y1")

    page["A2"] = "№"
    page["B2"] = "ФИО студента"
    page.merge_cells("C2:Y2")  # пустая объединённая шапка под даты

    for index, name in enumerate(names):
        row = 3 + index
        # Номера именно формулами, как в настоящем файле.
        page.cell(row=row, column=1, value=1 if index == 0 else f"=A{row - 1}+1")
        page.cell(row=row, column=2, value=name)

    note_row = 3 + len(names)
    page.cell(row=note_row, column=3, value="Староста: Тестовый Студент; контакт скрыт")
    page.merge_cells(start_row=note_row, start_column=3, end_row=note_row, end_column=25)

    workbook.save(path)
    return path


class Exhausted(Exception):
    """Данные в заглушке кончились — способ выйти из бесконечного цикла."""


class FakeSerial:
    """COM-порт без железа: отдаёт заранее заданный поток кусками
    указанного размера, а когда данные кончились — бросает Exhausted."""

    def __init__(self, data: bytes, chunk: int):
        self.data = data
        self.chunk = chunk
        self.pos = 0

    def read(self, _size: int) -> bytes:
        if self.pos >= len(self.data):
            raise Exhausted
        piece = self.data[self.pos : self.pos + self.chunk]
        self.pos += self.chunk
        return piece
