"""Тесты обратного направления: файлы групп → база.

Главное правило: база главнее файла. Импорт только дополняет её тем,
чего в ней нет, и никогда не перебивает отметку, сделанную картой.
"""

from datetime import date, datetime, time

import pytest
from openpyxl import load_workbook

from rfid import journal, roster as R
from rfid.codes import parse_line
from rfid.storage import IMPORT_RAW, CardConflict, MarkStatus, Storage
from test_roster import make_roster_file

CARD_A = parse_line("Em-Marine[A100] 007,42")
CARD_B = parse_line("Em-Marine[B200] 008,43")

GROUP = "1000000/10001"
NAMES = ["Тестов Тест Тестович", "Примеров Пример Примерович", "Образцова Проба"]
TODAY = date(2026, 9, 25)


@pytest.fixture
def db(tmp_path):
    with Storage(tmp_path / "test.db") as s:
        yield s


@pytest.fixture
def folder(tmp_path):
    subject = tmp_path / "tables" / "Бургеростроение 1 курс"
    subject.mkdir(parents=True)
    make_roster_file(subject / "10001.xlsx", names=NAMES, group=GROUP)
    return R.discover_subjects(tmp_path / "tables")[0]


def put(folder, cells: dict[tuple[int, int], object]) -> None:
    """Вписать в файл группы значения, как это сделал бы человек в Excel."""
    path = folder.path / "10001.xlsx"
    book = load_workbook(path)
    for (row, col), value in cells.items():
        book.active.cell(row=row, column=col, value=value)
    book.save(path)


def run(db, folder):
    [result] = journal.import_subject(db, folder, today=TODAY)
    assert result.ok, result.error
    return result


class TestReadingCells:
    @pytest.mark.parametrize("value, expected", [
        (time(10, 0), time(10, 0)),
        (datetime(2026, 9, 12, 9, 5, 30), time(9, 5)),
        ("09:05", time(9, 5)),
        ("9.05", time(9, 5)),
        (" 10:00:00 ", time(10, 0)),
        (None, None),
        ("", None),
        (R.ABSENT_MARK, None),
        ("-", None),
        ("н", None),
    ])
    def test_values(self, value, expected):
        assert R.cell_time(value) == expected

    @pytest.mark.parametrize("value", ["+", "был", "25:00", "10:75"])
    def test_unclear_is_not_guessed(self, value):
        assert R.cell_time(value) == value.strip()


class TestHeaderDates:
    def test_our_format_gets_this_year(self):
        assert R.header_date("12.09", TODAY) == date(2026, 9, 12)

    def test_future_means_last_year(self):
        """«12.09», прочитанное в январе, — прошлый сентябрь, а не будущий."""
        assert R.header_date("12.09", date(2027, 1, 15)) == date(2026, 9, 12)

    def test_full_date_and_excel_date(self):
        assert R.header_date("12.09.2025", TODAY) == date(2025, 9, 12)
        assert R.header_date(datetime(2026, 9, 12), TODAY) == date(2026, 9, 12)

    @pytest.mark.parametrize("value", ["Примечание", "31.02", None, 5])
    def test_not_a_date(self, value):
        assert R.header_date(value, TODAY) is None


class TestImport:
    def test_marks_from_file_land_in_db(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): time(10, 0), (4, 3): R.ABSENT_MARK})
        result = run(db, folder)

        assert result.added == 1
        assert result.dates == (date(2026, 9, 12),)
        subject = db.find_subject(folder.name)
        rows = {r.student.full_name: r for r in db.day_rows(date(2026, 9, 12), subject)}
        assert rows[NAMES[0]].at == datetime(2026, 9, 12, 10, 0)
        # В базе теперь весь список — видно и отсутствующих.
        assert not rows[NAMES[1]].present
        assert NAMES[2] in rows

    def test_people_are_created_without_cards(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        assert db.count_cards() == 0
        assert db.student_by_name(NAMES[0], GROUP).card_code is None

    def test_marks_are_labelled_as_imported(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        raw = db.conn.execute("SELECT raw FROM attendance").fetchone()["raw"]
        assert raw == f"{IMPORT_RAW}10001.xlsx"

    def test_second_import_changes_nothing(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): "10:00", (4, 3): "10:05"})
        assert run(db, folder).added == 2
        again = run(db, folder)
        assert (again.added, again.same) == (0, 2)
        assert db.conn.execute("SELECT COUNT(*) FROM attendance").fetchone()[0] == 2

    def test_database_wins(self, db, folder):
        """В обычном режиме другое время в файле базу не меняет."""
        subject = db.add_subject(folder.name)
        db.add_student(CARD_A, NAMES[0], GROUP)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 12, 9, 2))

        put(folder, {(2, 3): "12.09", (3, 3): "11:30"})
        result = run(db, folder)

        assert (result.added, result.updated, result.conflicts) == (0, 0, 1)
        row = next(r for r in db.day_rows(date(2026, 9, 12), subject)
                   if r.student.full_name == NAMES[0])
        assert row.at == datetime(2026, 9, 12, 9, 2)

    def test_same_minute_keeps_scan_seconds(self, db, folder):
        """В файле секунд нет — живой приход не должен «исправляться» на :00."""
        subject = db.add_subject(folder.name)
        db.add_student(CARD_A, NAMES[0], GROUP)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 12, 9, 2, 13))

        put(folder, {(2, 3): "12.09", (3, 3): "09:02"})
        assert run(db, folder).same == 1
        at = db.conn.execute("SELECT at FROM attendance").fetchone()["at"]
        assert at == "2026-09-12T09:02:13"

    @pytest.mark.parametrize("cell", [None, R.ABSENT_MARK, "н"])
    def test_absence_in_file_does_not_delete(self, db, folder, cell):
        """Прочерк мог поставить сам экспорт раньше прихода — по нему не удаляем."""
        subject = db.add_subject(folder.name)
        db.add_student(CARD_A, NAMES[0], GROUP)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 12, 9, 2))

        put(folder, {(2, 3): "12.09", (3, 3): cell})
        run(db, folder)
        assert db.conn.execute("SELECT COUNT(*) FROM attendance").fetchone()[0] == 1

    def test_file_fills_the_gaps(self, db, folder):
        """Нет строки в базе — это нехватка сведений, а не «не был»."""
        subject = db.add_subject(folder.name)
        db.add_student(CARD_A, NAMES[0], GROUP)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 12, 9, 2))

        put(folder, {(2, 3): "12.09", (3, 3): "09:02", (4, 3): "09:10"})
        assert run(db, folder).added == 1

    def test_unclear_cell_is_reported_not_imported(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): "+", (4, 3): "10:00"})
        result = run(db, folder)
        assert result.added == 1
        assert [value for _, value in result.skipped] == ["+"]

    def test_foreign_columns_are_ignored(self, db, folder):
        put(folder, {(2, 3): "Примечание", (3, 3): "10:00"})
        assert run(db, folder).added == 0

    def test_matches_existing_person_despite_spelling(self, db, folder):
        """Регистр и «ё» не плодят второго человека."""
        db.add_student(CARD_A, NAMES[0].upper(), GROUP)
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        assert len([s for s in db.list_students() if s.group_name == GROUP]) == 3
        assert db.count_cards() == 1

    def test_export_after_import_keeps_file_marks(self, db, folder):
        """Круг файл → база → файл ничего не теряет."""
        put(folder, {(2, 3): "12.09", (3, 3): time(10, 0)})
        run(db, folder)
        journal.fill_subject(db, db.find_subject(folder.name), folder)

        sheet = load_workbook(folder.path / "10001.xlsx").active
        assert sheet.cell(row=2, column=3).value == "12.09"
        assert sheet.cell(row=3, column=3).value == "10:00"
        assert sheet.cell(row=4, column=3).value == R.ABSENT_MARK

    def test_unclear_cell_survives_export(self, db, folder):
        """«+» импорт не понял — значит и затирать его прочерком нельзя."""
        put(folder, {(2, 3): "12.09", (3, 3): "10:00", (4, 3): "+"})
        run(db, folder)
        journal.fill_subject(db, db.find_subject(folder.name), folder)

        sheet = load_workbook(folder.path / "10001.xlsx").active
        assert sheet.cell(row=4, column=3).value == "+"
        assert sheet.cell(row=5, column=3).value == R.ABSENT_MARK

    def test_hand_edit_after_import_survives_export(self, db, folder):
        """Запись в файл сперва забирает дописанное руками, потом пишет.

        Дописанный опоздавший остаётся. Исправленное время — нет: главнее база.
        """
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        subject = db.find_subject(folder.name)
        journal.fill_subject(db, subject, folder)

        # Опоздавшего дописали руками, а другому исправили время.
        put(folder, {(4, 3): "10:07", (3, 3): "09:55"})
        journal.fill_subject(db, subject, folder)

        sheet = load_workbook(folder.path / "10001.xlsx").active
        assert sheet.cell(row=3, column=3).value == "10:00"
        assert sheet.cell(row=4, column=3).value == "10:07"
        assert db.conn.execute("SELECT COUNT(*) FROM attendance").fetchone()[0] == 2

    def test_file_not_rewritten_if_import_failed(self, db, folder, monkeypatch):
        """Не удалось забрать отметки из файла — значит и перезаписывать его нельзя."""
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        path = folder.path / "10001.xlsx"
        before = path.read_bytes()

        def broken(*_a, **_k):
            raise OSError("диск отвалился")
        monkeypatch.setattr(db, "import_mark", broken)

        [result] = journal.fill_subject(db, db.get_or_create_subject(folder.name), folder)
        assert not result.ok
        assert path.read_bytes() == before

    def test_broken_file_does_not_block_others(self, db, folder):
        (folder.path / "битый.xlsx").write_bytes(b"not a workbook")
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        folder = R.as_folder(folder.path)

        results = journal.import_subject(db, folder, today=TODAY)
        assert sum(r.added for r in results) == 1
        assert [r.ok for r in results].count(False) == 1


class TestCardAfterImport:
    """Карту привязали уже после того, как отметки втянули из таблицы."""

    def test_tap_on_imported_day_is_a_duplicate(self, db, folder):
        put(folder, {(2, 3): "25.09", (3, 3): "09:00"})
        run(db, folder)
        subject = db.find_subject(folder.name)

        db.add_student(CARD_A, NAMES[0], GROUP)
        tap = db.mark(CARD_A, subject, at=datetime(2026, 9, 25, 9, 40))
        assert tap.status is MarkStatus.DUPLICATE
        assert tap.first_at == datetime(2026, 9, 25, 9, 0)

    def test_binding_reuses_the_person(self, db, folder):
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        before = db.student_by_name(NAMES[0], GROUP)

        student = db.add_student(CARD_A, NAMES[0], GROUP)
        assert student.id == before.id
        assert db.find_student(CARD_A).id == before.id

    def test_scan_before_binding_wins_over_import(self, db, folder):
        """Живой приход и строка из таблицы за одно занятие — остаётся приход."""
        put(folder, {(2, 3): "25.09", (3, 3): "09:30"})
        run(db, folder)
        subject = db.find_subject(folder.name)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 25, 9, 5))  # пока ничья

        db.add_student(CARD_A, NAMES[0], GROUP)
        rows = db.conn.execute("SELECT at FROM attendance").fetchall()
        assert [r["at"] for r in rows] == ["2026-09-25T09:05:00"]

    def test_second_card_for_same_person_is_refused(self, db, folder):
        db.add_student(CARD_A, NAMES[0], GROUP)
        with pytest.raises(CardConflict):
            db.add_student(CARD_B, NAMES[0], GROUP)
        assert db.find_student(CARD_B) is None


class TestUnbindingKeepsImportedMarks:
    """Снятие привязки не должно стирать то, что пришло из таблиц."""

    def test_remove_card_keeps_person_and_file_marks(self, db, folder):
        db.add_student(CARD_A, NAMES[0], GROUP)
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)
        subject = db.find_subject(folder.name)
        db.mark(CARD_A, subject, at=datetime(2026, 9, 19, 9, 0))

        assert db.remove_student(CARD_A) is True

        person = db.student_by_name(NAMES[0], GROUP)
        assert person is not None and person.card_code is None
        # Из таблицы — остаётся за человеком.
        assert db.has_mark(person, subject, date(2026, 9, 12))
        # Картой — снова неизвестная: привязка могла быть ошибочной.
        assert [u.card_code for u in db.unknown_cards()] == [CARD_A.canonical]

    def test_unbind_all_keeps_people(self, db, folder):
        db.add_student(CARD_A, NAMES[0], GROUP)
        put(folder, {(2, 3): "12.09", (3, 3): "10:00"})
        run(db, folder)

        db.unbind_all_cards()
        assert db.count_cards() == 0
        assert db.count_students() == 3
        orphans = db.conn.execute(
            "SELECT COUNT(*) FROM attendance WHERE student_id IS NULL AND card_code IS NULL"
        ).fetchone()[0]
        assert orphans == 0


def age_file(folder, when: datetime) -> None:
    """Сделать вид, что файл последний раз сохраняли в when."""
    import os
    os.utime(folder.path / "10001.xlsx", (when.timestamp(), when.timestamp()))


def count(db) -> int:
    return db.conn.execute("SELECT COUNT(*) FROM attendance").fetchone()[0]


class TestFullSync:
    """Полная синхронизация — явное действие, в нём главнее таблицы."""

    def _scanned(self, db, folder, at):
        subject = db.get_or_create_subject(folder.name)
        db.add_student(CARD_A, NAMES[0], GROUP)
        db.mark(CARD_A, subject, at=at)
        return subject

    def test_time_from_file_wins(self, db, folder):
        subject = self._scanned(db, folder, datetime(2026, 9, 12, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): "11:30"})

        [result] = journal.import_subject(db, folder, TODAY, sync=True)
        assert result.updated == 1
        assert db.marks_of(db.find_student(CARD_A), subject, date(2026, 9, 12))[0][1] \
            == datetime(2026, 9, 12, 11, 30)
        assert count(db) == 1  # правится время, а не заводится вторая строка

    def test_dash_removes_mark(self, db, folder):
        self._scanned(db, folder, datetime(2026, 9, 12, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): R.ABSENT_MARK, (4, 3): "10:00"})
        age_file(folder, datetime(2026, 9, 13, 12, 0))

        [result] = journal.import_subject(db, folder, TODAY, sync=True)
        assert [(r.full_name, r.day) for r in result.removed] == [(NAMES[0], date(2026, 9, 12))]
        assert count(db) == 1  # осталась только отметка из файла

    def test_mark_newer_than_file_is_kept(self, db, folder):
        """Прочерк поставил прошлый экспорт, а карту приложили позже — не удалять."""
        put(folder, {(2, 3): "12.09", (3, 3): R.ABSENT_MARK})
        age_file(folder, datetime(2026, 9, 12, 9, 0))
        self._scanned(db, folder, datetime(2026, 9, 12, 9, 40))

        [result] = journal.import_subject(db, folder, TODAY, sync=True)
        assert result.removed == () and result.newer == 1
        assert count(db) == 1

    def test_unclear_cell_removes_nothing(self, db, folder):
        self._scanned(db, folder, datetime(2026, 9, 12, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): "+"})
        age_file(folder, datetime(2026, 9, 13, 12, 0))
        journal.import_subject(db, folder, TODAY, sync=True)
        assert count(db) == 1

    def test_days_absent_from_file_are_untouched(self, db, folder):
        """Колонки нет — таблица про этот день ничего не утверждает."""
        self._scanned(db, folder, datetime(2026, 9, 19, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): R.ABSENT_MARK})
        age_file(folder, datetime(2026, 9, 20, 12, 0))
        journal.import_subject(db, folder, TODAY, sync=True)
        assert count(db) == 1

    def test_dry_run_changes_nothing(self, db, folder):
        self._scanned(db, folder, datetime(2026, 9, 12, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): R.ABSENT_MARK, (4, 3): "10:00"})
        age_file(folder, datetime(2026, 9, 13, 12, 0))
        people = db.count_students()

        [plan] = journal.import_subject(db, folder, TODAY, sync=True, dry_run=True)
        assert len(plan.removed) == 1 and plan.added == 1
        assert count(db) == 1 and db.count_students() == people
        at = db.conn.execute("SELECT at FROM attendance").fetchone()["at"]
        assert at == "2026-09-12T09:02:00"

    def test_after_sync_files_and_db_agree(self, db, folder):
        self._scanned(db, folder, datetime(2026, 9, 19, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): "10:00", (4, 3): "10:05"})

        imported, filled = journal.sync_subject(db, folder, TODAY)
        assert all(r.ok for r in imported + filled)

        sheet = load_workbook(folder.path / "10001.xlsx").active
        # День только из базы дописан в файл, дни из файла — в базе.
        assert sheet.cell(row=2, column=4).value == "19.09"
        assert sheet.cell(row=3, column=4).value == "09:02"
        assert count(db) == 3


class TestSyncCommand:
    def _args(self, tmp_path, **kw):
        from types import SimpleNamespace
        base = dict(db=tmp_path / "cli.db", tables=tmp_path / "tables",
                    subject=None, sync=True, yes=False, choose=False)
        base.update(kw)
        return SimpleNamespace(**base)

    def _prepare(self, tmp_path, folder):
        with Storage(tmp_path / "cli.db") as s:
            subject = s.add_subject(folder.name)
            s.add_student(CARD_A, NAMES[0], GROUP)
            s.mark(CARD_A, subject, at=datetime(2026, 9, 12, 9, 2))
        put(folder, {(2, 3): "12.09", (3, 3): R.ABSENT_MARK})
        age_file(folder, datetime(2026, 9, 13, 12, 0))

    def test_declining_changes_nothing(self, tmp_path, folder, monkeypatch):
        from rfid import cli
        self._prepare(tmp_path, folder)
        before = (folder.path / "10001.xlsx").read_bytes()
        monkeypatch.setattr(cli, "_ask", lambda _p: "n")

        cli.cmd_import(self._args(tmp_path))
        with Storage(tmp_path / "cli.db") as s:
            assert count(s) == 1
        assert (folder.path / "10001.xlsx").read_bytes() == before

    def test_confirming_removes(self, tmp_path, folder, monkeypatch, capsys):
        from rfid import cli
        self._prepare(tmp_path, folder)
        monkeypatch.setattr(cli, "_ask", lambda _p: "y")

        assert cli.cmd_import(self._args(tmp_path)) == 0
        assert "Будут УДАЛЕНЫ" in capsys.readouterr().out
        with Storage(tmp_path / "cli.db") as s:
            assert count(s) == 0

    def test_menu_asks_the_mode(self, tmp_path, folder, monkeypatch):
        from rfid import cli
        self._prepare(tmp_path, folder)
        answers = iter(["1"])
        monkeypatch.setattr(cli, "_ask", lambda _p: next(answers))

        cli.cmd_import(self._args(tmp_path, sync=False, choose=True))
        with Storage(tmp_path / "cli.db") as s:
            assert count(s) == 1  # обычный режим ничего не удаляет
