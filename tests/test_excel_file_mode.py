"""Excel-файлы не должны оставаться открытыми после чтения или записи."""

from datetime import date, datetime

from openpyxl import load_workbook as real_load_workbook

from rfid import excel
from rfid.excel import attendance as attendance_module
from rfid.excel import roster as roster_module
from tests.helpers import make_roster_file


class WorkbookSpy:
    def __init__(self, workbook, closed: list[bool]):
        self.workbook = workbook
        self.closed = closed

    def __getattr__(self, name):
        return getattr(self.workbook, name)

    def __getitem__(self, name):
        return self.workbook[name]

    def close(self):
        self.closed.append(True)
        self.workbook.close()


def tracking_loader(calls: list[dict], closed: list[bool]):
    def load(*args, **kwargs):
        calls.append(kwargs)
        return WorkbookSpy(real_load_workbook(*args, **kwargs), closed)

    return load


def test_roster_is_read_in_read_only_mode_and_closed(tmp_path, monkeypatch):
    path = make_roster_file(tmp_path / "group.xlsx")
    calls: list[dict] = []
    closed: list[bool] = []
    monkeypatch.setattr(roster_module, "load_workbook", tracking_loader(calls, closed))

    excel.read_roster(path)

    assert calls == [{"read_only": True}]
    assert closed == [True]


def test_attendance_is_read_in_read_only_mode_and_closed(tmp_path, monkeypatch):
    path = make_roster_file(tmp_path / "group.xlsx")
    roster = excel.read_roster(path)
    calls: list[dict] = []
    closed: list[bool] = []
    monkeypatch.setattr(
        attendance_module,
        "load_workbook",
        tracking_loader(calls, closed),
    )

    excel.read_attendance(roster)

    assert calls == [{"read_only": True}]
    assert closed == [True]


def test_writable_book_is_closed_immediately_after_export(tmp_path, monkeypatch):
    path = make_roster_file(tmp_path / "group.xlsx")
    roster = excel.read_roster(path)
    calls: list[dict] = []
    closed: list[bool] = []
    monkeypatch.setattr(
        attendance_module,
        "load_workbook",
        tracking_loader(calls, closed),
    )

    excel.write_attendance(
        roster,
        {date(2026, 10, 16): {roster.students[0].full_name: datetime(2026, 10, 16, 9)}},
    )

    assert calls == [{"read_only": False}]
    assert closed == [True]
