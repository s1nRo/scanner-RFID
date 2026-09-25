"""Отметки в файле группы: колонки с датами справа от ФИО.

Запись сохраняет оформление файла и не затирает того, что вписано руками.
Чтение возвращает найденное как есть, а непонятное не угадывает.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..lessons import is_late, session_starts
from ..names import normalize_name
from .cells import ABSENT_MARK, DATE_HEADER_FORMAT, TIME_FORMAT, cell_time, header_date
from .roster import Roster, RosterStudent
from .save import save_atomically

LATE_FILL = PatternFill(fill_type="solid", fgColor="FFFFFF00")

_DATE_COLUMN_WIDTH = 7


# -------------------------------------------------------------------- запись


def write_attendance(
    roster: Roster,
    marks: dict[date, dict[str, datetime]],
    *,
    absent_mark: str = ABSENT_MARK,
) -> Path:
    """Проставить отметки в файле группы, сохранив его оформление.

    marks: дата -> {ФИО: время прихода}, общий на предмет, со всеми группами.
    ФИО сверяются без учёта регистра и «ё». Имена, которых нет в списке,
    молча не теряются — их возвращает сверка через unmatched_names().
    """
    workbook = load_workbook(roster.path)
    sheet = workbook[roster.sheet_title]

    _unmerge_header(sheet, roster.header_row, roster.first_date_col)
    columns = _date_columns(sheet, roster)

    for day in sorted(marks):
        column = columns.get(day.strftime(DATE_HEADER_FORMAT))
        if column is None:
            column = _add_date_column(sheet, roster, columns, day)

        # Приходы сначала режутся на занятия — см. lessons.session_starts().
        starts = session_starts(marks[day].values())
        present = {normalize_name(name): at for name, at in marks[day].items()}
        for student in roster.students:
            at = present.get(normalize_name(student.full_name))
            _free_cell(sheet, student.row, column)
            cell = sheet.cell(row=student.row, column=column)
            assert isinstance(cell, Cell), "объединение с ячейки снято выше"
            if at is None and cell_time(cell.value) is not None:
                # В файле что-то стоит, а в базе отметки нет: время, вписанное
                # руками, или непонятное «+». Файл главнее — не затираем.
                continue
            cell.value = at.strftime(TIME_FORMAT) if at else absent_mark
            cell.alignment = Alignment(horizontal="center")
            if at is not None and is_late(at, starts):
                cell.fill = LATE_FILL
            elif cell.fill == LATE_FILL:
                # Повторный экспорт должен снимать устаревшую подсветку.
                cell.fill = PatternFill()

    save_atomically(workbook, roster.path)
    return roster.path


def _unmerge_header(sheet: Worksheet, row: int, from_col: int) -> None:
    """Снять объединения, мешающие писать в строку заголовка дат.

    В настоящем списке C2:Y2 объединено — без этого дату туда не положить.
    """
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row <= row <= merged.max_row and merged.max_col >= from_col:
            sheet.unmerge_cells(str(merged))


def _free_cell(sheet: Worksheet, row: int, col: int) -> None:
    """Снять объединение, накрывающее ячейку.

    Например, ФИО в B17:C17 занимает будущую ячейку отметки C17.
    Освобождается только она; значение в B17 сохраняется.
    """
    for merged in list(sheet.merged_cells.ranges):
        if (merged.min_row <= row <= merged.max_row
                and merged.min_col <= col <= merged.max_col):
            sheet.unmerge_cells(str(merged))


def _date_columns(sheet: Worksheet, roster: Roster) -> dict[str, int]:
    """Какие даты уже есть в шапке и в каких колонках."""
    found: dict[str, int] = {}
    for col in range(roster.first_date_col, sheet.max_column + 1):
        value = sheet.cell(row=roster.header_row, column=col).value
        if isinstance(value, str) and value.strip():
            found[value.strip()] = col
        elif isinstance(value, datetime):
            found[value.strftime(DATE_HEADER_FORMAT)] = col
    return found


def _add_date_column(
    sheet: Worksheet, roster: Roster, columns: dict[str, int], day: date
) -> int:
    """Завести колонку для нового дня правее всех занятых."""
    title = day.strftime(DATE_HEADER_FORMAT)
    column = max([*columns.values(), roster.first_date_col - 1]) + 1
    columns[title] = column
    header = sheet.cell(row=roster.header_row, column=column, value=title)
    header.alignment = Alignment(horizontal="center")
    header.font = Font(bold=True)
    sheet.column_dimensions[get_column_letter(column)].width = _DATE_COLUMN_WIDTH
    return column


# -------------------------------------------------------------------- чтение


@dataclass(frozen=True, slots=True)
class FileMarks:
    """Отметки, найденные в файле группы."""

    # дата -> [(студент, время прихода)]
    marks: dict[date, list[tuple[RosterStudent, datetime]]]
    # Что не удалось понять: (заголовок колонки или ФИО, значение).
    # Такое не угадываем, а показываем человеку.
    skipped: list[tuple[str, str]]
    # дата -> кто в этот день явно не был (пусто или прочерк)
    absent: dict[date, list[RosterStudent]] = field(default_factory=dict)


def read_attendance(roster: Roster, today: date | None = None) -> FileMarks:
    """Прочитать уже проставленные в файле отметки."""
    today = today or date.today()
    sheet = load_workbook(roster.path)[roster.sheet_title]

    marks: dict[date, list[tuple[RosterStudent, datetime]]] = {}
    absent: dict[date, list[RosterStudent]] = {}
    skipped: list[tuple[str, str]] = []
    for col in range(roster.first_date_col, sheet.max_column + 1):
        head = sheet.cell(row=roster.header_row, column=col).value
        day = header_date(head, today)
        if day is None:
            continue  # не дата — чужая колонка, например «Примечание»
        for student in roster.students:
            value = cell_time(sheet.cell(row=student.row, column=col).value)
            if value is None:
                absent.setdefault(day, []).append(student)
            elif isinstance(value, str):
                skipped.append((f"{head}, {student.full_name}", value))
            else:
                marks.setdefault(day, []).append((student, datetime.combine(day, value)))
    return FileMarks(marks, skipped, absent)
