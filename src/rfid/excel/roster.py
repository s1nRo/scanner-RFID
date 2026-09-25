"""Разбор присланного списка группы: где шапка, кто в списке, какая группа.

Пример структуры списка с вымышленным номером группы:

    A1  «Группа 1000000/10001», объединено A1:Y1
    A2  «№», B2 «ФИО студента», C2:Y2 объединено и пусто
    A3+ номера — ФОРМУЛЫ вида =A2+1, а не числа
    B3+ фамилии
    C28:Y28 объединённое примечание «Староста: …» под списком

Поэтому: номера берём по порядку строк, а не из ячеек; конец списка ищем
по пустому ФИО. Колонки с датами идут сразу справа от ФИО.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..names import normalize_name

# Сколько строк сверху просматривать в поисках шапки.
_HEADER_SEARCH_DEPTH = 15
# Сколько колонок слева просматривать в поисках шапки и названия группы.
_SEARCH_WIDTH = 10

_GROUP_RE = re.compile(r"группа\s*[:№]?\s*(?P<name>\S.*)", re.IGNORECASE)
_NAME_HEADER = re.compile(r"ФИО|фамилия", re.IGNORECASE)

# Номер группы «по виду»: 1000000/10001, 1000000-10001 и подобные.
# Нужен, когда слова «Группа» в файле нет — например «ИКНК 1000000/10002».
_GROUP_CODE_RE = re.compile(r"\d{3,}\s*[/\\.\-]\s*\d{2,}")


class RosterError(Exception):
    """Файл не похож на список группы."""


@dataclass(frozen=True, slots=True)
class RosterStudent:
    number: int      # порядковый номер в списке
    full_name: str
    row: int         # строка на листе


@dataclass(slots=True)
class Roster:
    path: Path
    sheet_title: str
    group_name: str
    header_row: int
    name_col: int
    first_date_col: int
    students: list[RosterStudent] = field(default_factory=list)

    def find(self, full_name: str) -> RosterStudent | None:
        target = normalize_name(full_name)
        for student in self.students:
            if normalize_name(student.full_name) == target:
                return student
        return None


def read_roster(path: str | Path) -> Roster:
    path = Path(path)
    sheet = load_workbook(path).active
    if not isinstance(sheet, Worksheet):
        raise RosterError(f"в {path.name} активный лист — не таблица")

    header_row, name_col = _detect_header(sheet)
    students = _read_students(sheet, header_row, name_col)
    if not students:
        raise RosterError(f"в {path.name} не нашлось ни одного студента")

    return Roster(
        path=path,
        sheet_title=sheet.title,
        group_name=_detect_group(sheet, header_row, fallback=path.stem),
        header_row=header_row,
        name_col=name_col,
        first_date_col=name_col + 1,
        students=students,
    )


def unmatched_names(roster: Roster, names: list[str]) -> list[str]:
    """Кого из отмеченных не нашлось в списке группы.

    Связка «карта → студент» живёт в базе и опирается на ФИО, поэтому
    правка фамилии в файле рвёт связь. Молчать об этом нельзя.
    """
    return [name for name in names if roster.find(name) is None]


def _detect_header(sheet: Worksheet) -> tuple[int, int]:
    """Найти строку шапки и колонку с ФИО."""
    depth = min(sheet.max_row, _HEADER_SEARCH_DEPTH)
    for row in range(1, depth + 1):
        for col in range(1, min(sheet.max_column, _SEARCH_WIDTH) + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and _NAME_HEADER.search(value):
                return row, col
    raise RosterError("не нашёл шапку: нет ячейки со словом «ФИО»")


def _read_students(sheet: Worksheet, header_row: int, name_col: int) -> list[RosterStudent]:
    students: list[RosterStudent] = []
    row = header_row + 1
    while row <= sheet.max_row:
        value = sheet.cell(row=row, column=name_col).value
        if value is None or not str(value).strip():
            break  # список кончился; ниже могут быть примечания
        students.append(
            RosterStudent(number=len(students) + 1, full_name=str(value).strip(), row=row)
        )
        row += 1
    return students


def _cells_above_header(sheet: Worksheet, header_row: int):
    for row in range(1, header_row):
        for col in range(1, min(sheet.max_column, _SEARCH_WIDTH) + 1):
            value = sheet.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                yield value.strip()


def _detect_group(sheet: Worksheet, header_row: int, fallback: str) -> str:
    """Название группы из шапки файла.

    Три попытки по убыванию надёжности: явное «Группа X»; номер группы
    по виду (например, «ИКНК 1000000/10002» — без слова
    «Группа»); наконец первая непустая строка над шапкой. Имя файла —
    последнее средство, оно хуже всего описывает группу.
    """
    above = list(_cells_above_header(sheet, header_row))

    for text in above:
        match = _GROUP_RE.match(text)
        if match:
            return match.group("name").strip()

    for text in above:
        match = _GROUP_CODE_RE.search(text)
        if match:
            return re.sub(r"\s*", "", match.group(0))

    return above[0] if above else fallback
